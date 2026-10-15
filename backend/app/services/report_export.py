"""
Выгрузка отчёта конструктора в Excel (xlsx) и PDF.

На вход — результат run_report(): показатели, строки, итоги, сравнение.
  • xlsx: шапка, таблица с числовыми форматами (суммы, проценты), строка «Итого»,
    закреплённая шапка и автофильтр; на втором листе — график по первому показателю.
  • PDF: шапка, плитки итогов, график и таблица; кириллица — шрифтом DejaVu
    (в контейнере ставится пакетом fonts-dejavu-core).
"""
from __future__ import annotations

import io
import os
from datetime import date, datetime
from xml.sax.saxutils import escape

from app.core.utils import TZ_LOCAL

BRAND = "#328FCD"
GREEN = "#1A9E6E"
RED = "#E03A5A"
MUTED = "#6F9FBD"
TEXT = "#0D2236"
HAIR = "#D6E8F5"

GRAN_LABEL = {"day": "по дням", "week": "по неделям", "month": "по месяцам"}


# ─── Общие помощники ─────────────────────────────────────────────────

def _d(s: str | None) -> str:
    return date.fromisoformat(s).strftime("%d.%m.%Y") if s else ""


def period_line(rep: dict) -> str:
    line = f"Период: {_d(rep['date_from'])} – {_d(rep['date_to'])}"
    if rep.get("prev_from"):
        line += f"  ·  сравнение с {_d(rep['prev_from'])} – {_d(rep['prev_to'])}"
    return line


def _group(n: float) -> str:
    return f"{n:,.0f}".replace(",", " ")


def fmt_value(v, unit: str) -> str:
    if v is None:
        return "—"
    if unit == "rub":
        return f"{_group(v)} ₽"
    if unit == "pct":
        return f"{v:.1f}".replace(".", ",") + "%"
    if unit == "sec":
        v = int(round(v))
        return f"{v // 60} мин {v % 60:02d} с" if v >= 60 else f"{v} с"
    return _group(v)


def signed_delta(d: dict | None) -> float | None:
    if not d or d.get("pct") is None:
        return None
    return -d["pct"] if d.get("dir") == "down" else d["pct"]


def fmt_delta(d: dict | None) -> str:
    s = signed_delta(d)
    if s is None:
        return "—"
    return f"{'+' if s > 0 else '−' if s < 0 else ''}{abs(s)}%"


def delta_good(metric: dict, d: dict | None) -> bool | None:
    s = signed_delta(d)
    if s is None or s == 0 or not metric.get("better"):
        return None
    return (s > 0) == (metric["better"] == "up")


ROW_CMP_SHORT = {"day": "к пред. дню", "week": "к пред. неделе", "month": "к пред. месяцу"}


def delta_header(rep: dict) -> str:
    if rep.get("row_compare") == "bucket":
        return ROW_CMP_SHORT.get(rep["granularity"], "изм.")
    return "к прошл. периоду"


def compare_note(rep: dict) -> str:
    if not rep.get("prev_from"):
        return ""
    if rep.get("row_compare") == "bucket":
        unit = {"day": "предыдущим днём", "week": "предыдущей неделей", "month": "предыдущим месяцем"}[rep["granularity"]]
        return (f"Изменение в строках — по сравнению с {unit}; в итогах — с периодом "
                f"{_d(rep['prev_from'])} – {_d(rep['prev_to'])}.")
    return f"Изменение — по сравнению с периодом {_d(rep['prev_from'])} – {_d(rep['prev_to'])} той же длины."


def filters_line(rep: dict) -> str:
    parts = [f"{f['label']}: {', '.join(f['values'])}" for f in rep.get("filters") or []]
    return ("Фильтр — " + "; ".join(parts)) if parts else ""


def row_label(row: dict) -> str:
    return " · ".join(row["labels"])


def file_stem(title: str, rep: dict) -> str:
    safe = "".join(ch for ch in title if ch.isalnum() or ch in " -_").strip() or "Отчёт"
    return f"{safe} {rep['date_from']}–{rep['date_to']}"


def _now_local() -> str:
    return datetime.now(TZ_LOCAL).strftime("%d.%m.%Y %H:%M")


# ─── Excel ───────────────────────────────────────────────────────────────

XLSX_FORMATS = {
    "rub": '#,##0 "₽"',
    "count": "#,##0",
    "min": "#,##0",
    "sec": '0 "с"',
    "pct": '0.0"%"',
}


def to_xlsx(rep: dict, title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"

    metrics = rep["metrics"]
    dims = rep["dimensions"]
    compare = rep.get("prev_from") is not None

    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True, color=TEXT[1:])
    ws["A2"] = f"{rep['source']['label']} · {period_line(rep)}"
    ws["A2"].font = Font(color=MUTED[1:])
    ws["A3"] = f"Выгружено из «Пульса» {_now_local()}"
    ws["A3"].font = Font(color=MUTED[1:], size=9)
    if filters_line(rep):
        ws["A4"] = filters_line(rep)
        ws["A4"].font = Font(color=BRAND[1:], bold=True)

    head = [d["label"] for d in dims] or ["Показатель"]
    cols: list[tuple[str, dict, str]] = []   # (kind, metric, header)
    for m in metrics:
        cols.append(("cur", m, m["label"]))
        if compare:
            prev_name = "пред." if rep.get("row_compare") == "bucket" else "прошлый период"
            cols.append(("prev", m, f"{m['label']}, {prev_name}"))
            cols.append(("delta", m, f"{m['label']}, {delta_header(rep)}"))
    header_row = 5
    thin = Side(style="thin", color=HAIR[1:])
    for i, h in enumerate(head + [c[2] for c in cols], start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BRAND[1:])
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if i > len(head) else "left")
    ws.row_dimensions[header_row].height = 44

    def write_values(r: int, values: dict, prev: dict | None, delta: dict, bold=False):
        col = len(head) + 1
        for kind, m, _h in cols:
            if kind == "cur":
                v = values.get(m["key"])
                fmt = XLSX_FORMATS.get(m["unit"], "#,##0")
            elif kind == "prev":
                v = (prev or {}).get(m["key"])
                fmt = XLSX_FORMATS.get(m["unit"], "#,##0")
            else:
                s = signed_delta(delta.get(m["key"]))
                v = s / 100 if s is not None else None
                fmt = '+0%;-0%;0%'
            cell = ws.cell(row=r, column=col, value=v)
            cell.number_format = fmt
            cell.border = Border(bottom=thin)
            if bold:
                cell.font = Font(bold=True)
            if kind == "delta":
                good = delta_good(m, delta.get(m["key"]))
                if good is not None:
                    cell.font = Font(bold=bold, color=(GREEN if good else RED)[1:])
            col += 1

    r = header_row + 1
    first_data = r
    for row in rep["rows"]:
        for i, lab in enumerate(row["labels"], start=1):
            c = ws.cell(row=r, column=i, value=lab)
            c.border = Border(bottom=thin)
        write_values(r, row["values"], row.get("prev"), row.get("delta") or {})
        r += 1
    last_data = r - 1

    tot = rep["totals"]
    ws.cell(row=r, column=1, value="Итого").font = Font(bold=True)
    for i in range(1, len(head) + 1):
        ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor="E6F3FB")
    write_values(r, tot["values"], tot.get("prev"), tot.get("delta") or {}, bold=True)
    for i in range(len(head) + 1, len(head) + len(cols) + 1):
        ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor="E6F3FB")

    for i in range(1, len(head) + 1):
        width = max([len(str(ws.cell(row=x, column=i).value or "")) for x in range(header_row, r + 1)] + [10])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
    for i in range(len(head) + 1, len(head) + len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = ws.cell(row=header_row + 1, column=len(head) + 1)
    if rep["rows"]:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(head) + len(cols))}{last_data}"
    note = compare_note(rep)
    if note:
        ws.cell(row=r + 2, column=1, value=note).font = Font(color=MUTED[1:], size=9)
    if rep.get("truncated"):
        ws.cell(row=r + 3, column=1, value="Показаны первые 500 строк — сузьте отчёт фильтрами").font = Font(color=RED[1:])

    # График по первому показателю (для одной группировки)
    if len(dims) == 1 and rep["rows"]:
        m = metrics[0]
        is_period = dims[0]["key"] == "period"
        chart = LineChart() if is_period and len(rep["rows"]) > 12 else BarChart()
        if isinstance(chart, BarChart):
            chart.type = "col" if is_period or len(rep["rows"]) <= 12 else "bar"
            chart.gapWidth = 60
        chart.title = m["label"]
        chart.height, chart.width = 9, 22
        chart.legend = None
        data = Reference(ws, min_col=len(head) + 1, min_row=header_row, max_row=min(last_data, first_data + 59))
        cats = Reference(ws, min_col=1, min_row=first_data, max_row=min(last_data, first_data + 59))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        try:
            chart.series[0].graphicalProperties.line.solidFill = BRAND[1:]
            chart.series[0].graphicalProperties.solidFill = BRAND[1:]
        except Exception:  # noqa: BLE001 — оформление графика не критично
            pass
        cs = wb.create_sheet("График")
        cs["A1"] = f"{title}: {m['label']}"
        cs["A1"].font = Font(size=13, bold=True)
        cs.add_chart(chart, "A3")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── PDF ─────────────────────────────────────────────────────────────────

FONT_CANDIDATES = [
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"),
]
_font_ready: tuple[str, str] | None = None


def _fonts() -> tuple[str, str]:
    """Регистрирует шрифт с кириллицей; если его нет — PDF будет без русских букв, поэтому ошибка."""
    global _font_ready
    if _font_ready:
        return _font_ready
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for regular, bold in FONT_CANDIDATES:
        if os.path.exists(regular):
            pdfmetrics.registerFont(TTFont("Pulse", regular))
            pdfmetrics.registerFont(TTFont("Pulse-Bold", bold if os.path.exists(bold) else regular))
            from reportlab.lib.fonts import addMapping
            addMapping("Pulse", 0, 0, "Pulse")
            addMapping("Pulse", 1, 0, "Pulse-Bold")
            _font_ready = ("Pulse", "Pulse-Bold")
            return _font_ready
    raise RuntimeError("Не найден шрифт DejaVu для PDF (пакет fonts-dejavu-core)")


def _chart(rep: dict, width: float, font: str):
    """Столбики (или линия для длинных рядов) по первому показателю; для периода — с прошлым периодом."""
    from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors

    dims = rep["dimensions"]
    if not dims or not rep["rows"]:
        return None
    m = rep["metrics"][0]
    if len(dims) == 2 and any(d["key"] == "period" for d in dims):
        return _stacked_chart(rep, width, font)
    rows = rep["rows"]
    is_period = len(dims) == 1 and dims[0]["key"] in ("period", "weekday", "hour")
    if not is_period:
        rows = sorted(rows, key=lambda r: -(r["values"].get(m["key"]) or 0))[:15]
    vals = [float(r["values"].get(m["key"]) or 0) for r in rows]
    ps = rep.get("prev_series")
    prev = [float(x.get(m["key"]) or 0) for x in ps] if ps and dims[0]["key"] == "period" else None
    labels = [" · ".join(r["labels"]) for r in rows]

    def short(lbl: str) -> str:
        if dims[0]["key"] == "period" and rep["granularity"] in ("day", "week"):
            return lbl[:5]
        return lbl if len(lbl) <= 28 else lbl[:27] + "…"

    axis_fmt = _axis_fmt

    brand, grey = colors.HexColor(BRAND), colors.HexColor("#C5D9E8")
    if is_period:
        h = 190
        d = Drawing(width, h)
        n = len(vals)
        if n > 40:
            c = HorizontalLineChart()
            c.lines[0].strokeColor = brand
            c.lines[0].strokeWidth = 1.6
            if prev:
                c.lines[1].strokeColor = grey
                c.lines[1].strokeWidth = 1.2
        else:
            c = VerticalBarChart()
            c.bars[0].fillColor = brand
            c.bars[0].strokeColor = None
            if prev:
                c.bars[1].fillColor = grey
                c.bars[1].strokeColor = None
            c.barSpacing = 0.5
            c.groupSpacing = 3
        c.x, c.y, c.width, c.height = 48, 30, width - 60, h - 52
        c.data = [vals] + ([prev] if prev else [])
        step = max(1, n // 14)
        c.categoryAxis.categoryNames = [short(x) if i % step == 0 else "" for i, x in enumerate(labels)]
        c.categoryAxis.labels.fontName = font
        c.categoryAxis.labels.fontSize = 6.5
        c.categoryAxis.labels.angle = 0 if n <= 14 else 45
        c.categoryAxis.labels.boxAnchor = "n" if n <= 14 else "ne"
        c.categoryAxis.strokeColor = colors.HexColor(HAIR)
        c.valueAxis.labels.fontName = font
        c.valueAxis.labels.fontSize = 7
        c.valueAxis.labelTextFormat = axis_fmt
        c.valueAxis.valueMin = 0
        c.valueAxis.strokeColor = colors.HexColor(HAIR)
        c.valueAxis.gridStrokeColor = colors.HexColor("#E8F2F9")
        c.valueAxis.visibleGrid = True
        d.add(c)
        d.add(String(48, h - 12, m["label"], fontName=font, fontSize=8.5, fillColor=colors.HexColor(TEXT)))
        if prev:
            d.add(String(width - 160, h - 12, "■", fontName=font, fontSize=9, fillColor=brand))
            d.add(String(width - 150, h - 12, "этот период", fontName=font, fontSize=7.5, fillColor=colors.HexColor(MUTED)))
            d.add(String(width - 92, h - 12, "■", fontName=font, fontSize=9, fillColor=grey))
            d.add(String(width - 82, h - 12, "прошлый", fontName=font, fontSize=7.5, fillColor=colors.HexColor(MUTED)))
        return d

    n = len(vals)
    h = 34 + n * 16
    d = Drawing(width, h)
    c = HorizontalBarChart()
    c.x, c.y, c.width, c.height = 170, 8, width - 230, h - 30
    c.data = [list(reversed(vals))]
    c.bars[0].fillColor = brand
    c.bars[0].strokeColor = None
    c.barSpacing = 1
    c.groupSpacing = 4
    c.categoryAxis.categoryNames = [short(x) for x in reversed(labels)]
    c.categoryAxis.labels.fontName = font
    c.categoryAxis.labels.fontSize = 7
    c.categoryAxis.labels.boxAnchor = "e"
    c.categoryAxis.labels.dx = -4
    c.categoryAxis.strokeColor = colors.HexColor(HAIR)
    c.valueAxis.labels.fontName = font
    c.valueAxis.labels.fontSize = 7
    c.valueAxis.labelTextFormat = axis_fmt
    c.valueAxis.valueMin = 0
    c.valueAxis.strokeColor = colors.HexColor(HAIR)
    c.valueAxis.visibleGrid = True
    c.valueAxis.gridStrokeColor = colors.HexColor("#E8F2F9")
    c.barLabelFormat = lambda v: fmt_value(v, m["unit"])
    c.barLabels.fontName = font
    c.barLabels.fontSize = 6.5
    c.barLabels.boxAnchor = "w"
    c.barLabels.dx = 3
    c.barLabels.fillColor = colors.HexColor(TEXT)
    d.add(c)
    title = m["label"] + (" — первые 15" if len(rep["rows"]) > 15 else "")
    d.add(String(0, h - 12, title, fontName=font, fontSize=8.5, fillColor=colors.HexColor(TEXT)))
    return d


SERIES_COLORS = ["#328FCD", "#1A9E6E", "#C98418", "#8B5CF6", "#0A9FB8", "#DB2777"]
OTHER_COLOR = "#94B3C9"
ADDITIVE_UNITS = {"count", "rub", "min"}


def _axis_fmt(v):
    a = abs(v)
    if a >= 1_000_000:
        return f"{v / 1_000_000:.1f}".replace(".", ",") + " млн"
    if a >= 10_000:
        return f"{v / 1000:.0f} тыс"
    return _group(v)


def _stacked_chart(rep: dict, width: float, font: str):
    """Период × другая группировка: столбики по периодам, сложенные из шести крупнейших значений."""
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors
    from reportlab.pdfbase.pdfmetrics import stringWidth

    m = rep["metrics"][0]
    dims = rep["dimensions"]
    pi = 0 if dims[0]["key"] == "period" else 1
    oi = 1 - pi
    periods: dict[str, str] = {}
    for r in sorted(rep["rows"], key=lambda r: r["keys"][pi]):
        periods.setdefault(r["keys"][pi], r["labels"][pi])
    totals: dict[str, float] = {}
    for r in rep["rows"]:
        totals[r["labels"][oi]] = totals.get(r["labels"][oi], 0) + float(r["values"].get(m["key"]) or 0)
    top = [k for k, _ in sorted(totals.items(), key=lambda kv: -kv[1])[:len(SERIES_COLORS)]]
    additive = m["unit"] in ADDITIVE_UNITS
    series = top + (["Остальные"] if additive and len(totals) > len(top) else [])
    pkeys = list(periods)
    data = {sname: [0.0] * len(pkeys) for sname in series}
    for r in rep["rows"]:
        name = r["labels"][oi] if r["labels"][oi] in top else "Остальные"
        if name in data:
            data[name][pkeys.index(r["keys"][pi])] += float(r["values"].get(m["key"]) or 0)

    def color(i, name):
        return colors.HexColor(OTHER_COLOR if name == "Остальные" else SERIES_COLORS[i])

    # легенда в строку(и)
    legend_rows, row, x = [], [], 0.0
    for i, name in enumerate(series):
        w = 14 + stringWidth(name, font, 7.5) + 12
        if x + w > width and row:
            legend_rows.append(row); row, x = [], 0.0
        row.append((i, name, x)); x += w
    legend_rows.append(row)
    leg_h = 12 * len(legend_rows)

    h = 200 + leg_h
    d = Drawing(width, h)
    d.add(String(0, h - 10, m["label"], fontName=font, fontSize=8.5, fillColor=colors.HexColor(TEXT)))
    for li, lrow in enumerate(legend_rows):
        y = h - 26 - li * 12
        for i, name, lx in lrow:
            d.add(Rect(lx, y - 1, 8, 8, fillColor=color(i, name), strokeColor=None))
            d.add(String(lx + 11, y, name, fontName=font, fontSize=7.5, fillColor=colors.HexColor(MUTED)))

    n = len(pkeys)
    c = VerticalBarChart() if additive else HorizontalLineChart()
    c.x, c.y, c.width, c.height = 48, 30, width - 60, h - 52 - leg_h
    c.data = [data[sname] for sname in series]
    if additive:
        c.categoryAxis.style = "stacked"
        c.barSpacing = 0
        c.groupSpacing = 6 if n > 12 else 14
        for i, name in enumerate(series):
            c.bars[i].fillColor = color(i, name)
            c.bars[i].strokeColor = colors.white
            c.bars[i].strokeWidth = 0.5
    else:
        for i, name in enumerate(series):
            c.lines[i].strokeColor = color(i, name)
            c.lines[i].strokeWidth = 1.6
    step = max(1, n // 14)
    labels = [periods[k] for k in pkeys]
    if rep["granularity"] in ("day", "week"):
        labels = [lbl[:5] for lbl in labels]
    c.categoryAxis.categoryNames = [lbl if i % step == 0 else "" for i, lbl in enumerate(labels)]
    c.categoryAxis.labels.fontName = font
    c.categoryAxis.labels.fontSize = 6.5
    c.categoryAxis.labels.angle = 0 if n <= 14 else 45
    c.categoryAxis.labels.boxAnchor = "n" if n <= 14 else "ne"
    c.categoryAxis.strokeColor = colors.HexColor(HAIR)
    c.valueAxis.labels.fontName = font
    c.valueAxis.labels.fontSize = 7
    c.valueAxis.labelTextFormat = _axis_fmt
    c.valueAxis.valueMin = 0
    c.valueAxis.strokeColor = colors.HexColor(HAIR)
    c.valueAxis.visibleGrid = True
    c.valueAxis.gridStrokeColor = colors.HexColor("#E8F2F9")
    d.add(c)
    return d


def to_pdf(rep: dict, title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold = _fonts()
    metrics = rep["metrics"]
    dims = rep["dimensions"]
    compare = rep.get("prev_from") is not None
    n_cols = max(1, len(dims)) + len(metrics) * (2 if compare else 1)
    pagesize = landscape(A4) if n_cols > 6 else A4
    margin = 14 * mm
    width = pagesize[0] - 2 * margin

    st_title = ParagraphStyle("t", fontName=bold, fontSize=16, leading=19, textColor=colors.HexColor(TEXT))
    st_sub = ParagraphStyle("s", fontName=font, fontSize=8.5, leading=11, textColor=colors.HexColor(MUTED))
    st_cell = ParagraphStyle("c", fontName=font, fontSize=7.5, leading=9.5, textColor=colors.HexColor(TEXT))
    st_cell_b = ParagraphStyle("cb", parent=st_cell, fontName=bold)
    st_num = ParagraphStyle("n", parent=st_cell, alignment=TA_RIGHT)
    st_num_b = ParagraphStyle("nb", parent=st_num, fontName=bold)
    st_head = ParagraphStyle("h", parent=st_cell, fontName=bold, fontSize=7, leading=8.5, textColor=colors.white)
    st_head_r = ParagraphStyle("hr", parent=st_head, alignment=TA_RIGHT)
    st_kpi_l = ParagraphStyle("kl", fontName=font, fontSize=7, leading=9, textColor=colors.HexColor(MUTED))
    st_kpi_v = ParagraphStyle("kv", fontName=bold, fontSize=13, leading=16, textColor=colors.HexColor(TEXT))

    story = [
        Paragraph(escape(title), st_title),
        Spacer(1, 3),
        Paragraph(escape(f"{rep['source']['label']} · {period_line(rep)}")
                  + (f" · {GRAN_LABEL.get(rep['granularity'], '')}" if any(d['key'] == 'period' for d in dims) else ""),
                  st_sub),
        Paragraph(f"Выгружено из «Пульса» {_now_local()}", st_sub),
    ]
    if filters_line(rep):
        story.append(Paragraph(f'<font color="{BRAND}">{escape(filters_line(rep))}</font>', st_sub))
    story.append(Spacer(1, 10))

    # Плитки итогов (до 4 в ряд)
    tot = rep["totals"]

    def delta_html(m, d):
        txt = fmt_delta(d)
        good = delta_good(m, d)
        color = GREEN if good else RED if good is False else MUTED
        return f'<font color="{color}">{txt}</font>' if txt != "—" else ""

    tiles = []
    for m in metrics:
        v = fmt_value(tot["values"].get(m["key"]), m["unit"])
        sub = delta_html(m, (tot.get("delta") or {}).get(m["key"])) if compare else ""
        tiles.append([Paragraph(escape(m["label"].upper()), st_kpi_l), Paragraph(escape(v), st_kpi_v),
                      Paragraph(sub + (" к прошлому периоду" if sub else ""), st_kpi_l)])
    per_row = 4
    for i in range(0, len(tiles), per_row):
        chunk = tiles[i:i + per_row]
        cw = width / per_row
        t = Table([chunk], colWidths=[cw] * len(chunk), hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0, colors.white),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F1F7FC")),
            ("LINEAFTER", (0, 0), (-2, -1), 3, colors.white),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ]))
        story += [t, Spacer(1, 4)]
    story.append(Spacer(1, 8))

    chart = _chart(rep, width, font)
    if chart is not None:
        story += [chart, Spacer(1, 10)]

    # Таблица
    if rep["rows"]:
        head = [Paragraph(escape(d["label"]), st_head) for d in dims]
        for m in metrics:
            head.append(Paragraph(escape(m["label"]), st_head_r))
            if compare:
                head.append(Paragraph(delta_header(rep), st_head_r))
        data = [head]

        def value_cells(values, delta, b=False):
            out = []
            for m in metrics:
                out.append(Paragraph(escape(fmt_value(values.get(m["key"]), m["unit"])), st_num_b if b else st_num))
                if compare:
                    d = (delta or {}).get(m["key"])
                    good = delta_good(m, d)
                    color = GREEN if good else RED if good is False else MUTED
                    out.append(Paragraph(f'<font color="{color}">{fmt_delta(d)}</font>', st_num_b if b else st_num))
            return out

        prev_labels: list[str] = []
        for row in rep["rows"]:
            # повтор первой группировки не печатаем — так читается как сгруппированная таблица
            shown = [("" if i == 0 and len(dims) > 1 and prev_labels and prev_labels[0] == lbl else lbl)
                     for i, lbl in enumerate(row["labels"])]
            prev_labels = row["labels"]
            data.append([Paragraph(escape(lbl), st_cell_b if i == 0 and len(dims) > 1 else st_cell) for i, lbl in enumerate(shown)]
                        + value_cells(row["values"], row.get("delta")))
        total_row = [Paragraph("Итого", st_cell_b)] + [Paragraph("", st_cell)] * (len(dims) - 1)
        data.append(total_row + value_cells(tot["values"], tot.get("delta"), b=True))

        # ширина числовых колонок — по самому длинному значению (жирный итог шире)
        from reportlab.pdfbase.pdfmetrics import stringWidth
        num_w = []
        for m in metrics:
            texts = [fmt_value(r["values"].get(m["key"]), m["unit"]) for r in rep["rows"]]
            longest_word = max(m["label"].split(), key=len)
            w = max([stringWidth(t, font, 7.5) for t in texts]
                    + [stringWidth(fmt_value(tot["values"].get(m["key"]), m["unit"]), bold, 7.5),
                       stringWidth(longest_word, bold, 7)])
            num_w.append(max(44.0, w + 12))
            if compare:
                num_w.append(max(40.0, max(stringWidth(x, bold, 7) for x in delta_header(rep).split()) + 12))
        label_w = max(80.0, (width - sum(num_w)) / max(1, len(dims)))
        col_widths = [label_w] * len(dims) + num_w
        scale = width / sum(col_widths)
        if scale < 1:
            col_widths = [w * scale for w in col_widths]
        t = LongTable(data, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BRAND)),
            ("LINEBELOW", (0, 1), (-1, -2), 0.4, colors.HexColor(HAIR)),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E6F3FB")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t)
        if compare_note(rep):
            story += [Spacer(1, 4), Paragraph(compare_note(rep), st_sub)]
        if rep.get("truncated"):
            story += [Spacer(1, 4), Paragraph("Показаны первые 500 строк — сузьте отчёт фильтрами.", st_sub)]

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 7)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawString(margin, 8 * mm, "Пульс — аналитика Labvita")
        canvas.drawRightString(pagesize[0] - margin, 8 * mm, f"стр. {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=pagesize, leftMargin=margin, rightMargin=margin,
                            topMargin=margin, bottomMargin=16 * mm, title=title, author="Пульс")
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()
